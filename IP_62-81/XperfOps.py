# Write a Python Program to perform different operation on excel file

from openpyxl import Workbook, load_workbook


wb=load_workbook("sample.xlsx")
ws=wb.active

ws["F1"]="Total"
ws["G1"]="Average"

for row in range(2,ws.max_row+1):
    maths=ws[f"C{row}"].value
    science=ws[f"D{row}"].value
    english=ws[f"E{row}"].value

    total=maths+science+english
    avg=total/3

    ws[f"F{row}"]=total
    ws[f"G{row}"]=avg

wb.save("sample.xlsx")
print("operations performed succesfully")





































# wb=Workbook()
# ws=wb.active

# data = [
#     ["Roll No", "Name", "Maths", "Science", "English"],
#     [101, "Aarav", 85, 78, 90],
#     [102, "Vivaan", 72, 88, 76],
#     [103, "Aditya", 91, 82, 89],
#     [104, "Krishna", 65, 70, 68],
#     [105, "Ishaan", 88, 92, 84],
#     [106, "Arjun", 77, 75, 80],
#     [107, "Rohan", 93, 89, 95],
#     [108, "Karan", 58, 62, 64],
#     [109, "Rahul", 81, 85, 79],
#     [110, "Amit", 74, 69, 72],
#     [111, "Sohan", 87, 91, 86],
#     [112, "Mohan", 66, 73, 70],
#     [113, "Deepak", 95, 94, 97],
#     [114, "Vikas", 79, 81, 77],
#     [115, "Ankit", 83, 76, 88],
#     [116, "Yash", 90, 87, 92],
#     [117, "Nitin", 71, 68, 74],
#     [118, "Ritesh", 64, 66, 61],
#     [119, "Pankaj", 86, 84, 82],
#     [120, "Harsh", 92, 90, 94]
# ]

# for row in data:
#     ws.append(row)

# wb.save("sample.xlsx")



# wb=load_workbook("sample.xlsx")
# ws=wb.active

# for row in ws.iter_rows(values_only=True):
#     print(row)