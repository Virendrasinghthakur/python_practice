# Write a Python Program to write data to excel cell using python other than xlsxwriter library

import openpyxl as ox

wb = ox.Workbook()

ws = wb.active

ws["A1"] = "Name"
ws["B1"] = "Marks"

ws["A2"] = "Virendra"
ws["B2"] = 95

wb.save("student.xlsx")

print("Data written successfully.")