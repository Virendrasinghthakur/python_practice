# Write a Python Program to Create a excel file using python using other than xlswriter library

from openpyxl import Workbook

# Create a workbook
wb = Workbook()

# Select the active worksheet
ws = wb.active

# Write data
ws['A1'] = "Name"
ws['B1'] = "Marks"

ws['A2'] = "Virendra"
ws['B2'] = 95

ws['A3'] = "Rahul"
ws['B3'] = 88

# Save the workbook
wb.save("student.xlsx")

print("Excel file created successfully.")