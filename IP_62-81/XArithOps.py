# Write a Python Program to perform arithmetic operation on excel file.

import openpyxl as ox

wb=ox.load_workbook("sample.xlsx")
ws=wb.active

ws["H1"]="Percentage"

for row in range(2,ws.max_row+1):
    total=ws[f"F{row}"].value

    per=total/3

    ws[f"H{row}"]=per


for row in ws.iter_rows(values_only=True):
    print(row)


wb.save("sample.xlsx")