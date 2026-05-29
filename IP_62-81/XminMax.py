# Write a Python Program to find maximum and minimum number from a excel file

import openpyxl as ox
wb=ox.load_workbook("sample.xlsx")
ws=wb.active

maxp=float('-inf')
maxst=""
minp=float('inf')
minst=""
for row in range(2,ws.max_row+1):
    per=ws[f"H{row}"].value
    if per>maxp:
        maxp=per
        maxst=ws[f"B{row}"].value

    if per<minp:
        minp=per
        minst=ws[f"B{row}"].value


print(f"Student with maximum percentag is {maxst} with {maxp:.2f}%")
print(f"Student with minimum percentag is {minst} with {minp:.2f}%")