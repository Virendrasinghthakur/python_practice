# Write a Python Program to copy data from one excel sheet to another

import openpyxl as ox

wb1=ox.load_workbook("sample.xlsx")
ws1=wb1.active

wb2=ox.Workbook()
ws2=wb2.active

# print(ws1)
for row in ws1.iter_rows(values_only=True):
    ws2.append(row)

# print(ws2)
for row in ws2.iter_rows(values_only=True):
    print(row)


wb2.save("destination.xlsx")